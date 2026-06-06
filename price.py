import pandas as pd
import re
import os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

folder = "."
output_file = "combined_written_prices.xlsx"

yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

def process_file(file_name):
    df = pd.read_excel(file_name, header=None)

    products = []
    current_product = None
    current_prices = []

    for _, row in df.iterrows():
        first_cell = str(row[0]).strip() if pd.notna(row[0]) else ""

        if re.match(r"^[A-Z]{1,3}\d+\s*-", first_cell, re.IGNORECASE):
            if current_product and current_prices:
                products.append([current_product, current_prices])

            current_product = re.sub(
                r"^[A-Z]{1,3}\d+\s*-\s*",
                "",
                first_cell,
                flags=re.IGNORECASE
            ).strip()

            current_prices = []

        try:
            price = row[8]  # Column I = Cmimi
            if pd.notna(price):
                current_prices.append(float(price))
        except:
            pass

    if current_product and current_prices:
        products.append([current_product, current_prices])

    products = [p for p in products if "lei" not in p[0].lower()]

    result = []

    for product_name, prices in products:
        written_price = max(prices)
        unusual = False
        reasons = []

        # Flag any price decrease
        for i in range(1, len(prices)):
            if prices[i] < prices[i - 1]:
                unusual = True
                reasons.append("Price decrease")

        # Flag temporary change like 40,40,50,40,40
        for i in range(1, len(prices) - 1):
            if prices[i] != prices[i - 1] and prices[i] != prices[i + 1] and prices[i - 1] == prices[i + 1]:
                unusual = True
                reasons.append("Temporary price change")

        # Flag sudden final change like 40,40,40,50
        if len(prices) >= 3:
            if prices[-1] != prices[-2] and prices[:-1].count(prices[-2]) >= 2:
                unusual = True
                reasons.append("Sudden final change")

        result.append({
            "product": product_name,
            "price": written_price,
            "unusual": unusual,
            "reason": ", ".join(sorted(set(reasons)))
        })

    result.sort(key=lambda x: x["product"].lower())
    return result


excel_files = [
    f for f in os.listdir(folder)
    if f.endswith(".xlsx")
    and not f.startswith("~$")
    and f != output_file
]

wb = Workbook()
ws = wb.active
ws.title = "Written Prices"

current_row = 1

for file_name in excel_files:
    results = process_file(file_name)
    header_name = file_name.replace(".xlsx", "")

    ws.cell(row=current_row, column=1).value = header_name
    ws.cell(row=current_row, column=1).font = Font(bold=True, size=16)

    ws.cell(row=current_row + 1, column=1).value = "Product"
    ws.cell(row=current_row + 1, column=2).value = "Written Price"
    ws.cell(row=current_row + 1, column=3).value = "Flag Reason"

    for col in range(1, 4):
        ws.cell(row=current_row + 1, column=col).font = Font(bold=True)

    for row_num, item in enumerate(results, start=current_row + 2):
        ws.cell(row=row_num, column=1).value = item["product"]
        ws.cell(row=row_num, column=2).value = item["price"]
        ws.cell(row=row_num, column=3).value = item["reason"]

        if item["unusual"]:
            ws.cell(row=row_num, column=2).fill = yellow_fill

    current_row += len(results) + 4

# Auto-fit columns
for col in range(1, ws.max_column + 1):
    max_length = 0
    col_letter = get_column_letter(col)

    for cell in ws[col_letter]:
        if cell.value is not None:
            max_length = max(max_length, len(str(cell.value)))

    ws.column_dimensions[col_letter].width = min(max_length + 2, 45)

wb.save(output_file)

print(f"Done. Created {output_file}")