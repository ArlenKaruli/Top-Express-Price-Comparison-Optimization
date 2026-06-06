import pandas as pd
import re
import os
import sys
import tkinter as tk
from tkinter import filedialog
from collections import defaultdict
from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter


# =========================
# Let user choose Excel file
# =========================

root = tk.Tk()
root.withdraw()
root.attributes("-topmost", True)

selected_file = filedialog.askopenfilename(
    title="Zgjidh skedarin Excel",
    filetypes=[("Excel files", "*.xlsx *.xls")]
)

if not selected_file:
    print("Nuk u zgjodh asnjë skedar.")
    input("Shtyp Enter për të dalë...")
    sys.exit()

input_file = Path(selected_file)

# Save output in same folder as selected Excel file
timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
output_file = input_file.parent / f"prices_by_client_{timestamp}.xlsx"

print(f"Skedari i zgjedhur: {input_file.name}")


# =========================
# Settings
# =========================

yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")


# =========================
# Helper functions
# =========================

def clean_product_name(product_header):
    return re.sub(
        r"^[A-Z]{1,4}\d+\s*-\s*",
        "",
        str(product_header),
        flags=re.IGNORECASE
    ).strip()


def clean_client_name(client):
    client = str(client).strip()

    if "/" in client:
        client = client.split("/")[0].strip()

    return client


def is_product_header(value):
    if pd.isna(value):
        return False

    value = str(value).strip()

    return re.match(r"^[A-Z]{1,4}\d+\s*-", value, re.IGNORECASE) is not None


def find_column(df, possible_names, fallback_index):
    possible_names = [name.lower().strip() for name in possible_names]

    for row_index in range(min(20, len(df))):
        for col_index in range(len(df.columns)):
            cell = df.iloc[row_index, col_index]

            if pd.isna(cell):
                continue

            cell_text = str(cell).lower().strip()

            if cell_text in possible_names:
                return col_index

    return fallback_index


def check_unusual_prices(prices):
    unusual = False
    reasons = []

    # Flag if there is only one registration
    if len(prices) == 1:
        return True, "Vetëm një regjistrim çmimi"

    # Flag if there are only two prices and they differ
    if len(prices) == 2:
        if prices[0] != prices[1]:
            return True, "Vetëm dy çmime të ndryshme"
        else:
            return False, ""

    # Flag any price decrease
    for i in range(1, len(prices)):
        if prices[i] < prices[i - 1]:
            unusual = True
            reasons.append("Ulja e çmimit")

    # Flag temporary price change like 40, 40, 50, 40, 40
    for i in range(1, len(prices) - 1):
        previous_price = prices[i - 1]
        current_price = prices[i]
        next_price = prices[i + 1]

        if (
            current_price != previous_price
            and current_price != next_price
            and previous_price == next_price
        ):
            unusual = True
            reasons.append("Ndryshim i përkohshëm çmimi")

    # Flag sudden final change like 40, 40, 40, 50
    if len(prices) >= 3:
        last_price = prices[-1]
        previous_price = prices[-2]

        if last_price != previous_price and prices[:-1].count(previous_price) >= 2:
            unusual = True
            reasons.append("Ndryshim i papritur në fund")

    return unusual, ", ".join(sorted(set(reasons)))


# =========================
# Read Excel
# =========================

df = pd.read_excel(input_file, header=None)

client_col = find_column(
    df,
    ["klient", "client", "description", "pershkrim", "përshkrim"],
    fallback_index=4
)

price_col = find_column(
    df,
    ["cmimi", "çmimi", "price"],
    fallback_index=8
)

print(f"Kolona e klientit: {client_col}")
print(f"Kolona e çmimit: {price_col}")


# =========================
# Organize data by client/product
# =========================

client_products = defaultdict(lambda: defaultdict(list))

current_product = None

for _, row in df.iterrows():
    first_cell = row[0]

    if is_product_header(first_cell):
        current_product = clean_product_name(first_cell)
        continue

    if current_product is None:
        continue

    # Skip LEI products
    if "lei" in current_product.lower():
        continue

    client_value = row[client_col] if client_col < len(row) else None
    price_value = row[price_col] if price_col < len(row) else None

    if pd.isna(client_value) or pd.isna(price_value):
        continue

    try:
        price = float(price_value)
    except:
        continue

    client_name = clean_client_name(client_value)

    if not client_name:
        continue

    client_products[client_name][current_product].append(price)


# =========================
# Create output workbook
# =========================

wb = Workbook()
ws = wb.active
ws.title = "Prices By Client"

current_row = 1

for client_name in sorted(client_products.keys(), key=lambda x: x.lower()):
    products = client_products[client_name]

    # Client header
    ws.merge_cells(
        start_row=current_row,
        start_column=1,
        end_row=current_row,
        end_column=3
    )

    ws.cell(row=current_row, column=1).value = client_name
    ws.cell(row=current_row, column=1).font = Font(bold=True, size=16)
    ws.cell(row=current_row, column=1).alignment = Alignment(horizontal="left")

    # Table headers
    header_row = current_row + 1

    ws.cell(row=header_row, column=1).value = "Produkti"
    ws.cell(row=header_row, column=2).value = "Çmimi i Regjistruar"
    ws.cell(row=header_row, column=3).value = "Arsyeja e Kontrollit"

    for col in range(1, 4):
        ws.cell(row=header_row, column=col).font = Font(bold=True)

    row_num = current_row + 2

    for product_name in sorted(products.keys(), key=lambda x: x.lower()):
        prices = products[product_name]

        written_price = max(prices)
        unusual, reason = check_unusual_prices(prices)

        ws.cell(row=row_num, column=1).value = product_name
        ws.cell(row=row_num, column=2).value = written_price
        ws.cell(row=row_num, column=3).value = reason

        if unusual:
            ws.cell(row=row_num, column=2).fill = yellow_fill

        row_num += 1

    current_row = row_num + 3


# =========================
# Auto-fit columns
# =========================

for col in range(1, ws.max_column + 1):
    col_letter = get_column_letter(col)
    max_length = 0

    for cell in ws[col_letter]:
        if cell.value is not None:
            max_length = max(max_length, len(str(cell.value)))

    ws.column_dimensions[col_letter].width = min(max_length + 2, 50)


# =========================
# Save and open output
# =========================

wb.save(output_file)

print(f"Procesi përfundoi me sukses. U krijua skedari: {output_file.name}")

try:
    os.startfile(output_file)
except Exception as e:
    print(f"Skedari u krijua, por nuk u hap automatikisht: {e}")
    input("Shtyp Enter për të dalë...")